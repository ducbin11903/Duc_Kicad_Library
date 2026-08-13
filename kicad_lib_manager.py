#!/usr/bin/env python3
"""
KiCad Library Manager
Build EXE: pyinstaller --onefile --windowed --name KiCad_Lib_Manager kicad_lib_manager.py
"""
import os, sys, re, subprocess, shutil, threading, queue as TQ
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

APP_TITLE  = "KiCad Library Manager"
APP_VER    = "3.0"
APP_AUTHOR = "Duc"
APP_YEAR   = "2026"
SHEET_LIB="Library"; SHEET_Q="Queue"
LIB_PFX="Duc"; DEF_CAT="Misc"
CONFIG_FILE="kicad_lib_manager.cfg"

LIB_ROOT=os.getcwd()

def set_root(path):
    global LIB_ROOT
    LIB_ROOT=os.path.abspath(path)
    try:
        with open(CONFIG_FILE,"w",encoding="utf-8") as f: f.write(LIB_ROOT)
    except Exception: pass

def looks_like_root(d):
    """A library root has at least symbols/ or parts.xlsx in it."""
    if not os.path.isdir(d): return False
    return (os.path.isdir(os.path.join(d,"symbols")) or
            os.path.isfile(os.path.join(d,"parts.xlsx")))

def auto_detect_root(start):
    """Search the start folder, then up to 4 parents, then one level of children."""
    d=os.path.abspath(start)
    # walk upwards
    for _ in range(5):
        if looks_like_root(d): return d
        parent=os.path.dirname(d)
        if parent==d: break
        d=parent
    # look one level down from the start folder
    try:
        for name in os.listdir(start):
            sub=os.path.join(start,name)
            if looks_like_root(sub): return sub
    except Exception: pass
    return None

def load_root():
    """Restore the saved root; otherwise auto-detect from the current folder."""
    global LIB_ROOT
    try:
        if os.path.exists(CONFIG_FILE):
            q=open(CONFIG_FILE,encoding="utf-8").read().strip()
            if q and looks_like_root(q):
                LIB_ROOT=q; return
    except Exception: pass
    found=auto_detect_root(LIB_ROOT)
    if found: LIB_ROOT=found

def P(*parts):    return os.path.join(LIB_ROOT,*parts)
def parts_file(): return P("parts.xlsx")
def sym_dir():    return P("symbols")
def fp_dir():
    """footprints/<Prefix>.pretty if present, else footprints/ itself."""
    sub=P("footprints",f"{LIB_PFX}.pretty")
    if os.path.isdir(sub): return sub
    for d in (P("footprints"),):
        if os.path.isdir(d):
            pretty=[x for x in os.listdir(d) if x.endswith(".pretty")]
            if len(pretty)==1: return os.path.join(d,pretty[0])
            return d
    return sub

def m3d_dir():
    """3dmodels/<Prefix>.3dshapes if present, else 3dmodels/ itself."""
    sub=P("3dmodels",f"{LIB_PFX}.3dshapes")
    if os.path.isdir(sub): return sub
    for d in (P("3dmodels"),):
        if os.path.isdir(d):
            shapes=[x for x in os.listdir(d) if x.endswith(".3dshapes")]
            if len(shapes)==1: return os.path.join(d,shapes[0])
            return d
    return sub
def temp_dir():   return P("temp")
PIN_LEN=2.54; PIN_SPC=2.54; BODY_W=10.16; FSYM=1.27

CATS=["Resistors","Capacitors","Inductors","Diodes","Transistors","MCUs","ICs",
      "Power","Sensors","Connectors","Crystals","Memory","RF","Analog",
      "Interface","Optocouplers","Misc"]
PTYPES=["IN","OUT","IO","PWR","GND","PWR_OUT","PASSIVE","NC","OC","OE","UNSPEC","TRI","CLK"]
PSIDES=["L","R","T","B"]
REFS=["U","R","C","L","Q","J","D","Y","SW","F","BT","TP","M","P"]

CATMAP={"resistors":"Resistors","res":"Resistors","r":"Resistors",
 "capacitors":"Capacitors","cap":"Capacitors","c":"Capacitors",
 "inductors":"Inductors","ind":"Inductors","l":"Inductors",
 "diodes":"Diodes","diode":"Diodes","d":"Diodes",
 "transistors":"Transistors","mosfet":"Transistors","q":"Transistors",
 "mcus":"MCUs","mcu":"MCUs","ic":"ICs","ics":"ICs","logic":"ICs",
 "power":"Power","ldo":"Power","dcdc":"Power","sensor":"Sensors","sensors":"Sensors",
 "connectors":"Connectors","connector":"Connectors","conn":"Connectors",
 "crystals":"Crystals","crystal":"Crystals","memory":"Memory","rf":"RF",
 "analog":"Analog","interface":"Interface","optocouplers":"Optocouplers",
 "misc":"Misc","other":"Misc"}
PTMAP={"in":"input","out":"output","io":"bidirectional","bidir":"bidirectional",
 "pwr":"power_in","power":"power_in","vdd":"power_in","vcc":"power_in","gnd":"power_in",
 "pwr_out":"power_out","oc":"open_collector","oe":"open_emitter","nc":"no_connect",
 "passive":"passive","p":"passive","unspec":"unspecified","tri":"tri_state",
 "clk":"clock","clock":"clock"}

# ── Palette: industrial workshop (charcoal + copper + moss) ───────────────────
BG        = "#E8E6E1"   # warm light background
PANEL     = "#F4F2EE"   # panel surface
DARK      = "#2B3A42"   # charcoal header
DARK2     = "#37474F"   # secondary dark
COPPER    = "#C0611F"   # primary accent
COPPER_D  = "#9E4E19"
MOSS      = "#5B8C3E"   # success accent
MOSS_D    = "#456B2F"
SLATE     = "#546E7A"   # neutral button
SLATE_D   = "#3E5259"
GRID      = "#9E9A93"   # cell border
GRID_HDR  = "#1F2C33"   # header cell border
TXT       = "#1C1C1C"
TXT_INV   = "#FFFFFF"
CELL_BG   = "#FFFFFF"
CELL_ALT  = "#F0EEE9"
CELL_SEL  = "#FFD54F"   # selected row highlight
ERR       = "#B71C1C"
WARN      = "#E65100"

def resource(name):
    """Locate a bundled resource, both when run as a script and from the exe."""
    base=getattr(sys,"_MEIPASS",None) or os.path.dirname(os.path.abspath(__file__))
    p=os.path.join(base,name)
    if os.path.exists(p): return p
    p2=os.path.join(os.getcwd(),name)
    return p2 if os.path.exists(p2) else None


def apply_icon(win):
    """Set the window icon from app.ico, falling back to app_logo.png."""
    ico=resource("app.ico")
    if ico:
        try: win.iconbitmap(ico); return True
        except Exception: pass
    png=resource("app_logo.png")
    if png:
        try:
            img=tk.PhotoImage(file=png)
            win.iconphoto(True,img)
            win._icon_ref=img          # keep a reference alive
            return True
        except Exception: pass
    return False


def norm_cat(r): return CATMAP.get(r.strip().lower(), r.strip().title() if r.strip() else DEF_CAT)
def sym_path(c): return os.path.join(sym_dir(),f"{LIB_PFX}-{c}.kicad_sym")
def find_easyeda():
    sc=os.path.dirname(sys.executable)
    n="easyeda2kicad.exe" if os.name=="nt" else "easyeda2kicad"
    c=os.path.join(sc,n)
    return c if os.path.isfile(c) else shutil.which("easyeda2kicad")

def open_path(p, create=False):
    """Open a file or folder in the system file browser."""
    p=os.path.abspath(p)
    if not os.path.exists(p):
        if create:
            try: os.makedirs(p,exist_ok=True)
            except Exception: return False
        else:
            messagebox.showwarning("Folder not found",
                "This path does not exist:\n\n"+p+
                "\n\nUse CHANGE at the top right to pick the correct library folder.")
            return False
    try:
        if os.name=="nt": os.startfile(p)
        elif sys.platform=="darwin": subprocess.call(["open",p])
        else: subprocess.call(["xdg-open",p])
        return True
    except Exception: return False

def reveal_file(p):
    """Open explorer with file selected."""
    p=os.path.abspath(p)
    if not os.path.exists(p): return False
    try:
        if os.name=="nt": subprocess.Popen(f'explorer /select,"{p}"')
        elif sys.platform=="darwin": subprocess.call(["open","-R",p])
        else: subprocess.call(["xdg-open",os.path.dirname(p)])
        return True
    except Exception: return False

# ── Excel ─────────────────────────────────────────────────────────────────────
THIN=Border(left=Side(style="thin",color="D0D0D0"),right=Side(style="thin",color="D0D0D0"),
            top=Side(style="thin",color="D0D0D0"),bottom=Side(style="thin",color="D0D0D0"))
XFN=Font(name="Arial",size=10); XAC=Alignment(horizontal="center",vertical="center")
XAL=Alignment(horizontal="left",vertical="center")
XGR=PatternFill("solid",fgColor="F2F2F2"); XOK=PatternFill("solid",fgColor="E2EFDA")
XNO=PatternFill("solid",fgColor="FFDDE0")

def excel_write_sheet(sheet, rows, has_status=False):
    if not os.path.exists(parts_file()): return False
    try:
        wb=load_workbook(parts_file())
        if sheet not in wb.sheetnames: return False
        ws=wb[sheet]
        for r in ws.iter_rows(min_row=2):
            for c in r: c.value=None
        for ri,row in enumerate(rows,2):
            for ci,v in enumerate(row,1):
                cell=ws.cell(row=ri,column=ci,value=v)
                cell.font=XFN; cell.border=THIN
                if has_status and ci in (5,6,7):
                    cell.alignment=XAC
                    cell.fill=XOK if str(v).lower() in ("v","yes","x✓") else XNO
                elif has_status and ci==8:
                    cell.alignment=XAC; cell.fill=XGR
                else:
                    cell.alignment=XAC if ci==1 else XAL; cell.fill=XGR
        wb.save(parts_file()); return True
    except Exception: return False

def excel_read_sheet(sheet, ncols):
    rows=[]
    if not os.path.exists(parts_file()): return rows
    try:
        wb=load_workbook(parts_file())
        if sheet not in wb.sheetnames: return rows
        ws=wb[sheet]
        for row in ws.iter_rows(min_row=2,values_only=True):
            if not row or not row[0]: continue
            vals=[str(v).strip() if v is not None else "" for v in row[:ncols]]
            while len(vals)<ncols: vals.append("")
            rows.append(vals)
    except Exception: pass
    return rows

def excel_append_lib(lcsc,name,cat,note,s,f,m):
    if not os.path.exists(parts_file()): return
    try:
        wb=load_workbook(parts_file())
        if SHEET_LIB not in wb.sheetnames: return
        ws=wb[SHEET_LIB]; nr=ws.max_row+1
        now=datetime.now().strftime("%Y-%m-%d %H:%M")
        for col,val in enumerate([lcsc,name,cat,note,"v" if s else "x",
                                  "v" if f else "x","v" if m else "x",now],1):
            cell=ws.cell(row=nr,column=col,value=val); cell.font=XFN; cell.border=THIN
            if col in(5,6,7): cell.alignment=XAC; cell.fill=XOK if val=="v" else XNO
            elif col==8: cell.alignment=XAC; cell.fill=XGR
            else: cell.alignment=XAL; cell.fill=XGR
        wb.save(parts_file())
    except Exception: pass

# ── KiCad ─────────────────────────────────────────────────────────────────────
def ensure_dirs():
    for d in [sym_dir(),fp_dir(),m3d_dir(),temp_dir()]: os.makedirs(d,exist_ok=True)
def ensure_sym(p):
    if not os.path.exists(p):
        with open(p,"w",encoding="utf-8") as f:
            f.write("(kicad_symbol_lib (version 20231120) (generator duc_lib)\n)\n")
def clear_temp():
    if os.path.exists(temp_dir()): shutil.rmtree(temp_dir(),ignore_errors=True)
    os.makedirs(temp_dir(),exist_ok=True)

def extract_blocks(fp):
    if not os.path.exists(fp): return []
    with open(fp,"r",encoding="utf-8") as f: c=f.read()
    out,pos=[],0
    while True:
        s=c.find('(symbol "',pos)
        if s==-1: break
        st,e=0,-1
        for i in range(s,len(c)):
            if c[i]=="(": st+=1
            elif c[i]==")":
                st-=1
                if st==0: e=i+1; break
        if e==-1: break
        b=c[s:e]; m=re.match(r'\(symbol\s+"([^"]+)"',b)
        if m:
            lm=re.search(r'\(property\s+"LCSC Part"\s+"([^"]+)"',b)
            out.append({"name":m.group(1),"lcsc":lm.group(1) if lm else None,"content":b})
        pos=e
    return out

def merge_block(block,sf):
    with open(sf,"r",encoding="utf-8") as f: c=f.read()
    block=re.sub(r'\(property\s+"Footprint"\s+"[^:]+:([^"]+)"',
                 rf'(property "Footprint" "{LIB_PFX}:\1"',block)
    last=c.rfind(")")
    if last!=-1:
        with open(sf,"w",encoding="utf-8") as f: f.write(c[:last]+"\n  "+block+"\n"+c[last:])

def sym_exists(sf,n):
    if not os.path.exists(sf): return False
    with open(sf,"r",encoding="utf-8") as f: return f'(symbol "{n}"' in f.read()

def gen_kicad_sym(name,ref,desc,pins):
    L=[p for p in pins if p["side"]=="L"]; R=[p for p in pins if p["side"]=="R"]
    T=[p for p in pins if p["side"]=="T"]; B=[p for p in pins if p["side"]=="B"]
    n=max(len(L),len(R),1); nt=max(len(T),len(B),0)
    bh=n*PIN_SPC+PIN_SPC; bw=max(BODY_W,nt*PIN_SPC+PIN_SPC if nt else BODY_W)
    hw,hh=bw/2,bh/2
    yp=lambda k:[((k-1)*PIN_SPC/2)-i*PIN_SPC for i in range(k)] if k else []
    xp=lambda k:[-(k-1)*PIN_SPC/2+i*PIN_SPC for i in range(k)] if k else []
    f=lambda v:f"{v:.4f}"; FS=f"{FSYM:.4f}"; PL=f"{PIN_LEN:.4f}"
    def pr(nm,val,x=0.,y=0.,h=False):
        hd=" (hide yes)" if h else ""
        return f'  (property "{nm}" "{val}" (at {f(x)} {f(y)} 0)\n    (effects (font (size {FS} {FS})){hd})\n  )'
    def ps(p,x,y,a):
        pt=PTMAP.get(p["type"],"unspecified"); pn=p["name"].replace('"','\\"')
        return (f'    (pin {pt} line (at {f(x)} {f(y)} {a}) (length {PL})\n'
                f'      (name "{pn}" (effects (font (size {FS} {FS}))))\n'
                f'      (number "{p["number"]}" (effects (font (size {FS} {FS}))))\n    )')
    ln=[f'(symbol "{name}"','  (pin_names (offset 1.016))','  (exclude_from_sim no)',
        '  (in_bom yes)','  (on_board yes)',
        pr("Reference",ref,0,hh+1.27),pr("Value",name,0,-(hh+1.27)),
        pr("Footprint","",h=True),pr("Datasheet","",h=True),pr("Description",desc,h=True),
        f'  (symbol "{name}_0_1"',
        f'    (rectangle (start {f(-hw)} {f(hh)}) (end {f(hw)} {f(-hh)})',
        '      (stroke (width 0) (type default))','      (fill (type background))','    )','  )',
        f'  (symbol "{name}_1_1"']
    for p,y in zip(L,yp(len(L))): ln.append(ps(p,-(hw+PIN_LEN),y,0))
    for p,y in zip(R,yp(len(R))): ln.append(ps(p,hw+PIN_LEN,y,180))
    for p,x in zip(T,xp(len(T))): ln.append(ps(p,x,hh+PIN_LEN,270))
    for p,x in zip(B,xp(len(B))): ln.append(ps(p,x,-(hh+PIN_LEN),90))
    ln+=["  )",")"]; return "\n".join(ln)


# ==============================================================================
# COMPONENT FILE RESOLVER
# Symbol  --Footprint property-->  .kicad_mod  --model path-->  .step / .wrl
# ==============================================================================
def find_symbol_block(name, lcsc, cat=None):
    """Locate a symbol by name or LCSC. Returns (sym_file, block_dict) or (None,None)."""
    files=[]
    if cat:
        c=sym_path(cat)
        if os.path.exists(c): files.append(c)
    if os.path.isdir(sym_dir()):
        for fn in sorted(os.listdir(sym_dir())):
            if fn.startswith(f"{LIB_PFX}-") and fn.endswith(".kicad_sym"):
                fp=os.path.join(sym_dir(),fn)
                if fp not in files: files.append(fp)
    nl=(name or "").strip().lower()
    ll=(lcsc or "").strip().lower()
    for fp in files:
        for b in extract_blocks(fp):
            if nl and b["name"].strip().lower()==nl: return fp,b
            if ll and (b["lcsc"] or "").strip().lower()==ll: return fp,b
    return None,None


def footprint_of(block):
    """Footprint file stem stored in the symbol, e.g. 'Duc:SOT-23_L2.9' -> 'SOT-23_L2.9'."""
    if not block: return ""
    m=re.search(r'\(property\s+"Footprint"\s+"([^"]*)"',block["content"])
    if not m: return ""
    val=m.group(1).strip()
    return val.split(":")[-1] if val else ""


def find_file(root, basename, exts):
    """Find a file by stem or full name anywhere under root."""
    if not basename or not os.path.isdir(root): return None
    stem=os.path.splitext(basename)[0].lower()
    for dirpath,_,names in os.walk(root):
        for n in names:
            if not n.lower().endswith(exts): continue
            if os.path.splitext(n)[0].lower()==stem: return os.path.join(dirpath,n)
    return None


def models_in_footprint(fp_path):
    """Every 3D model file name referenced inside a .kicad_mod."""
    out=[]
    try: c=open(fp_path,encoding="utf-8").read()
    except Exception: return out
    for m in re.finditer(r'\(model\s+"([^"]+)"',c):
        out.append(os.path.basename(m.group(1).replace("\\","/")))
    return out


def resolve_component(name, lcsc, cat):
    """Resolve every file that belongs to one component, following the real links."""
    res={"sym_file":None,"sym_name":None,"fp_name":"","fp_path":None,
         "model_name":"","model_path":None}
    sf,block=find_symbol_block(name,lcsc,cat)
    if sf:
        res["sym_file"]=sf; res["sym_name"]=block["name"]
        fp_stem=footprint_of(block)
        res["fp_name"]=fp_stem
        if fp_stem:
            res["fp_path"]=find_file(fp_dir(),fp_stem,(".kicad_mod",))
    # fall back to a name match when the symbol carries no footprint field
    if not res["fp_path"]:
        for key in (name,lcsc):
            hit=find_file(fp_dir(),key,(".kicad_mod",))
            if hit: res["fp_path"]=hit; res["fp_name"]=os.path.splitext(os.path.basename(hit))[0]; break
    # 3D model comes from inside the footprint file
    if res["fp_path"]:
        for mn in models_in_footprint(res["fp_path"]):
            hit=find_file(m3d_dir(),mn,(".step",".stp",".wrl"))
            if hit:
                res["model_name"]=os.path.basename(hit); res["model_path"]=hit; break
            res["model_name"]=res["model_name"] or mn
    if not res["model_path"] and res["fp_name"]:
        hit=find_file(m3d_dir(),res["fp_name"],(".step",".stp",".wrl"))
        if hit: res["model_path"]=hit; res["model_name"]=os.path.basename(hit)
    return res


# ==============================================================================
# DATA GRID — bordered, editable, filterable, pasteable
# ==============================================================================
class DataGrid(tk.Frame):
    RN_W = 40   # row-number column width

    def __init__(self, parent, cols, widths, combos=None, on_change=None,
                 filters=True, **kw):
        super().__init__(parent, bg=GRID, bd=1, relief="solid", **kw)
        self.cols=cols; self.widths=widths
        self.combos=combos or {}
        self.on_change=on_change
        self._rows=[]
        self._sel=set()
        self._has_filters=filters
        self._filters=[tk.StringVar() for _ in cols] if filters else None
        self._build()

    # ── Layout ────────────────────────────────────────────────────────────────
    def _build(self):
        # HEADER
        hdr=tk.Frame(self,bg=GRID_HDR); hdr.pack(fill="x")
        corner=tk.Frame(hdr,bg=GRID_HDR,width=self.RN_W,height=30)
        corner.pack(side="left",padx=(0,1),pady=(0,1)); corner.pack_propagate(False)
        tk.Label(corner,text="#",bg="#1A252B",fg="#8FA3AD",
                 font=("Segoe UI",9,"bold")).pack(fill="both",expand=True)
        for c,w in zip(self.cols,self.widths):
            cell=tk.Frame(hdr,bg=GRID_HDR,width=w,height=30)
            cell.pack(side="left",padx=(0,1),pady=(0,1)); cell.pack_propagate(False)
            tk.Label(cell,text=c.upper(),bg=DARK,fg=TXT_INV,
                     font=("Segoe UI",9,"bold"),anchor="center").pack(fill="both",expand=True)

        # FILTER ROW (visually distinct from data)
        if self._has_filters:
            fr=tk.Frame(self,bg=GRID); fr.pack(fill="x")
            lab=tk.Frame(fr,bg=GRID,width=self.RN_W,height=24)
            lab.pack(side="left",padx=(0,1),pady=(0,1)); lab.pack_propagate(False)
            tk.Label(lab,text="FIND",bg="#6B5B3E",fg="#FFE9B8",
                     font=("Segoe UI",7,"bold")).pack(fill="both",expand=True)
            for i,w in enumerate(self.widths):
                cell=tk.Frame(fr,bg=GRID,width=w,height=24)
                cell.pack(side="left",padx=(0,1),pady=(0,1)); cell.pack_propagate(False)
                e=tk.Entry(cell,textvariable=self._filters[i],font=("Segoe UI",9),
                           bg="#FFF6DC",fg="#4A3B18",relief="flat",bd=0,
                           insertbackground=COPPER)
                e.pack(fill="both",expand=True,padx=1,pady=1)
                self._filters[i].trace_add("write",lambda *a:self._apply_filter())

        # BODY
        body=tk.Frame(self,bg=GRID); body.pack(fill="both",expand=True)
        self.canvas=tk.Canvas(body,bg=GRID,highlightthickness=0)
        vsb=ttk.Scrollbar(body,orient="vertical",command=self.canvas.yview)
        self.inner=tk.Frame(self.canvas,bg=GRID)
        self.inner.bind("<Configure>",
            lambda e:self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.create_window((0,0),window=self.inner,anchor="nw")
        self.canvas.configure(yscrollcommand=vsb.set)
        self.canvas.pack(side="left",fill="both",expand=True)
        vsb.pack(side="right",fill="y")
        for w in (self.canvas,self.inner):
            w.bind("<MouseWheel>",
                   lambda e:self.canvas.yview_scroll(-1*(e.delta//120),"units"))
        self.bind_all("<Control-v>",self._paste_evt,add="+")

    # ── Rows ──────────────────────────────────────────────────────────────────
    def add_row(self, values=None):
        vals=list(values) if values else [""]*len(self.cols)
        while len(vals)<len(self.cols): vals.append("")
        idx=len(self._rows)
        bg=CELL_BG if idx%2==0 else CELL_ALT
        rf=tk.Frame(self.inner,bg=GRID); rf.pack(fill="x")

        rec={}   # closures resolve the live index through this record

        # Row-number selector
        nc=tk.Frame(rf,bg=GRID,width=self.RN_W,height=26)
        nc.pack(side="left",padx=(0,1),pady=(0,1)); nc.pack_propagate(False)
        num_lbl=tk.Label(nc,text=str(idx+1),bg="#D5D1C8",fg="#3A3A3A",
                         font=("Segoe UI",9,"bold"),cursor="hand2")
        num_lbl.pack(fill="both",expand=True)
        num_lbl.bind("<Button-1>",        lambda e,d=rec:self.select_row(self._idx_of(d)))
        num_lbl.bind("<Control-Button-1>",lambda e,d=rec:self.toggle_row(self._idx_of(d)))
        num_lbl.bind("<Shift-Button-1>",  lambda e,d=rec:self._sel_range(self._idx_of(d)))

        vars_,widgets,cells=[],[],[]
        for i,(c,w) in enumerate(zip(self.cols,self.widths)):
            cell=tk.Frame(rf,bg=GRID,width=w,height=26)
            cell.pack(side="left",padx=(0,1),pady=(0,1)); cell.pack_propagate(False)
            v=tk.StringVar(value=str(vals[i]))

            if c in self.combos:
                # tk.OptionMenu (classic widget) so the background can be recoloured
                wd=tk.OptionMenu(cell,v,*self.combos[c])
                wd.config(bg=bg,fg=TXT,font=("Segoe UI",9),relief="flat",bd=0,
                          highlightthickness=0,anchor="w",padx=4,pady=0,
                          activebackground=CELL_SEL,cursor="hand2")
                wd["menu"].config(bg=CELL_BG,fg=TXT,font=("Segoe UI",9),
                                  activebackground=COPPER,activeforeground=TXT_INV)
                wd.bind("<Button-1>",lambda e,d=rec:self.select_row(self._idx_of(d)),add="+")
            else:
                wd=tk.Entry(cell,textvariable=v,font=("Segoe UI",9),
                            bg=bg,fg=TXT,relief="flat",bd=0,
                            highlightthickness=1,highlightbackground=bg,
                            highlightcolor=COPPER,insertbackground=TXT)
                wd.bind("<Button-1>",lambda e,d=rec:self.select_row(self._idx_of(d)),add="+")
                wd.bind("<FocusIn>", lambda e,d=rec:self.select_row(self._idx_of(d)),add="+")
                wd.bind("<Tab>",   lambda e,d=rec,ci=i:self._next_cell(self._idx_of(d),ci))
                wd.bind("<Return>",lambda e,d=rec,ci=i:self._next_row(self._idx_of(d),ci))

            wd.pack(fill="both",expand=True,padx=1,pady=1)
            if self.on_change: v.trace_add("write",lambda *a:self.on_change())
            vars_.append(v); widgets.append(wd); cells.append(cell)

        rec.update({"vars":vars_,"frame":rf,"widgets":widgets,"cells":cells,
                    "num":num_lbl,"visible":True})
        self._rows.append(rec)
        return idx

    # ── Selection ─────────────────────────────────────────────────────────────
    def _idx_of(self, rec):
        """Current index of a row record (survives inserts and deletes)."""
        for i,r in enumerate(self._rows):
            if r is rec: return i
        return -1

    def select_row(self, idx):
        if idx<0: return
        self._sel={idx}; self._paint()

    def toggle_row(self, idx):
        if idx<0: return
        if idx in self._sel: self._sel.discard(idx)
        else: self._sel.add(idx)
        self._paint()

    def _sel_range(self, idx):
        """Shift+click: select from the first selected row to this one."""
        if idx<0: return
        if not self._sel: self._sel={idx}
        else:
            a=min(self._sel); b=idx
            self._sel=set(range(min(a,b),max(a,b)+1))
        self._paint()

    def select_all(self):
        self._sel=set(range(len(self._rows))); self._paint()

    def _paint(self):
        for i,r in enumerate(self._rows):
            on   = i in self._sel
            base = CELL_BG if i%2==0 else CELL_ALT
            bg   = CELL_SEL if on else base
            # row number
            r["num"].config(bg=COPPER if on else "#D5D1C8",
                            fg=TXT_INV if on else "#3A3A3A")
            # row + cell frames (fills the 1px gaps too)
            r["frame"].config(bg=COPPER if on else GRID)
            for cell in r.get("cells",[]):
                cell.config(bg=COPPER if on else GRID)
            # widgets
            for wd in r["widgets"]:
                try:
                    if isinstance(wd,tk.Entry):
                        wd.config(bg=bg,highlightbackground=bg)
                    else:                       # tk.OptionMenu
                        wd.config(bg=bg,activebackground=bg)
                except Exception:
                    pass

    def selected_rows(self):
        return [[v.get().strip() for v in self._rows[i]["vars"]]
                for i in sorted(self._sel) if i<len(self._rows)]

    def selected_index(self):
        return sorted(self._sel)[0] if self._sel else None

    # ── Navigation ────────────────────────────────────────────────────────────
    def _next_cell(self,r,ci):
        if r<0 or r>=len(self._rows): return "break"
        nxt=(ci+1)%len(self.cols)
        self._rows[r]["widgets"][nxt].focus_set(); return "break"

    def _next_row(self,r,ci):
        if r<0: return "break"
        if r+1<len(self._rows): self._rows[r+1]["widgets"][ci].focus_set()
        else:
            self.add_row(); self._rows[-1]["widgets"][ci].focus_set()
        return "break"

    # ── Mutations ─────────────────────────────────────────────────────────────
    def delete_selected(self):
        if not self._sel: return
        for i in sorted(self._sel,reverse=True):
            if i<len(self._rows):
                self._rows[i]["frame"].destroy(); self._rows.pop(i)
        self._sel.clear(); self._rebuild_indices()
        if self.on_change: self.on_change()

    def clear(self):
        for r in self._rows: r["frame"].destroy()
        self._rows.clear(); self._sel.clear()
        if self.on_change: self.on_change()

    def _rebuild_indices(self):
        """Renumber the row headers. Bindings resolve indices live, so they stay valid."""
        for i,r in enumerate(self._rows):
            r["num"].config(text=str(i+1))
        self._paint()

    def get_rows(self):
        return [[v.get().strip() for v in r["vars"]] for r in self._rows]

    def load(self, rows):
        self.clear()
        for row in rows: self.add_row(row)

    def count(self): return len(self._rows)

    # ── Filter ────────────────────────────────────────────────────────────────
    def _apply_filter(self):
        if not self._filters: return
        terms=[f.get().strip().lower() for f in self._filters]
        for r in self._rows:
            show=all((not t) or (t in r["vars"][i].get().lower())
                     for i,t in enumerate(terms))
            if show and not r["visible"]:
                r["frame"].pack(fill="x"); r["visible"]=True
            elif not show and r["visible"]:
                r["frame"].pack_forget(); r["visible"]=False

    def clear_filters(self):
        if self._filters:
            for f in self._filters: f.set("")

    # ── Paste ─────────────────────────────────────────────────────────────────
    def _paste_evt(self,event):
        w=self.focus_get()
        if not w: return
        p=w; inside=False
        while p is not None:
            if p is self: inside=True; break
            p=getattr(p,"master",None)
        if not inside: return
        self.paste_clipboard(); return "break"

    def paste_clipboard(self):
        try: clip=self.clipboard_get()
        except Exception: return
        lines=[l for l in clip.replace("\r","").split("\n") if l.strip()]
        for line in lines:
            cells=[c.strip() for c in (line.split("\t") if "\t" in line else line.split(","))]
            while len(cells)<len(self.cols): cells.append("")
            cells=cells[:len(self.cols)]
            for i,c in enumerate(self.cols):
                if c=="Category" and cells[i]: cells[i]=norm_cat(cells[i])
                elif c=="Type" and cells[i]:   cells[i]=cells[i].upper()
                elif c=="Side" and cells[i]:   cells[i]=cells[i].upper()[:1]
            self.add_row(cells)
        if self.on_change: self.on_change()

# ── Button ────────────────────────────────────────────────────────────────────
class Btn(tk.Button):
    def __init__(self,parent,text,cmd,color=SLATE,dark=SLATE_D,fg=TXT_INV,**kw):
        super().__init__(parent,text=text,command=cmd,bg=color,fg=fg,
                         font=("Segoe UI",9,"bold"),relief="flat",bd=0,
                         padx=12,pady=6,cursor="hand2",
                         activebackground=dark,activeforeground=fg,**kw)
        self.bind("<Enter>",lambda e:self.config(bg=dark))
        self.bind("<Leave>",lambda e:self.config(bg=color))

# ── Log ───────────────────────────────────────────────────────────────────────
class LogBox(tk.Frame):
    def __init__(self,parent,height=6,**kw):
        super().__init__(parent,bg=GRID,bd=1,relief="solid",**kw)
        self.txt=tk.Text(self,height=height,font=("Consolas",9),bg="#22282B",
                         fg="#D6D0C8",state="disabled",relief="flat",wrap="word")
        sb=ttk.Scrollbar(self,command=self.txt.yview)
        self.txt.configure(yscrollcommand=sb.set)
        self.txt.pack(side="left",fill="both",expand=True,padx=1,pady=1)
        sb.pack(side="right",fill="y")
        for t,c in [("ok","#8BC34A"),("err","#EF5350"),("warn","#FFB74D"),
                    ("info","#81D4FA"),("dim","#90A4AE")]:
            self.txt.tag_config(t,foreground=c)
    def log(self,m,t="info"):
        self.txt.configure(state="normal"); self.txt.insert("end",m+"\n",t)
        self.txt.see("end"); self.txt.configure(state="disabled")
    def clear(self):
        self.txt.configure(state="normal"); self.txt.delete("1.0","end")
        self.txt.configure(state="disabled")

# ==============================================================================
# QUEUE TAB
# ==============================================================================
class QueueTab(tk.Frame):
    COLS=["LCSC","Name","Category","Note"]; W=[110,210,130,300]
    def __init__(self,parent,status):
        super().__init__(parent,bg=BG)
        self._status=status; self._running=False
        self._build(); self._load()

    def _build(self):
        top=tk.Frame(self,bg=BG); top.pack(fill="x",padx=10,pady=(10,6))
        tk.Label(top,text="QUEUE IMPORT",font=("Segoe UI",13,"bold"),
                 bg=BG,fg=DARK).pack(side="left")
        tk.Label(top,text="Add LCSC part numbers, then import symbols, footprints and 3D models",
                 font=("Segoe UI",9),bg=BG,fg="#5D5952").pack(side="left",padx=14)

        # Quick add
        qa=tk.Frame(self,bg=PANEL,bd=1,relief="solid"); qa.pack(fill="x",padx=10,pady=(0,8))
        inner=tk.Frame(qa,bg=PANEL); inner.pack(fill="x",padx=10,pady=8)
        tk.Label(inner,text="LCSC:",font=("Segoe UI",9,"bold"),bg=PANEL,fg=TXT).pack(side="left")
        self.v_lcsc=tk.StringVar()
        e1=tk.Entry(inner,textvariable=self.v_lcsc,width=14,font=("Segoe UI",10),
                    relief="solid",bd=1); e1.pack(side="left",padx=(4,12))
        e1.bind("<Return>",lambda ev:self._quick_add())
        tk.Label(inner,text="Name:",font=("Segoe UI",9,"bold"),bg=PANEL,fg=TXT).pack(side="left")
        self.v_name=tk.StringVar()
        tk.Entry(inner,textvariable=self.v_name,width=20,font=("Segoe UI",10),
                 relief="solid",bd=1).pack(side="left",padx=(4,12))
        tk.Label(inner,text="Category:",font=("Segoe UI",9,"bold"),bg=PANEL,fg=TXT).pack(side="left")
        self.v_cat=tk.StringVar(value="ICs")
        ttk.Combobox(inner,textvariable=self.v_cat,values=CATS,width=14,
                     state="readonly").pack(side="left",padx=(4,12))
        Btn(inner,"ADD ROW",self._quick_add,COPPER,COPPER_D).pack(side="left")
        tk.Label(inner,text="Paste from clipboard with Ctrl+V",
                 font=("Segoe UI",9,"bold"),bg=PANEL,fg=COPPER).pack(side="right")

        # Grid
        self.grid_w=DataGrid(self,self.COLS,self.W,
                             combos={"Category":CATS},on_change=self._upd)
        self.grid_w.pack(fill="both",expand=True,padx=10)

        # Toolbar
        tb=tk.Frame(self,bg=BG); tb.pack(fill="x",padx=10,pady=6)
        Btn(tb,"+ ROW",self._add_row,SLATE,SLATE_D).pack(side="left",padx=(0,4))
        Btn(tb,"DELETE",self.grid_w.delete_selected,"#8D6E63","#6D4C41").pack(side="left",padx=4)
        Btn(tb,"CLEAR",self.grid_w.clear,"#78909C","#546E7A").pack(side="left",padx=4)
        Btn(tb,"PASTE",self.grid_w.paste_clipboard,SLATE,SLATE_D).pack(side="left",padx=4)
        tk.Frame(tb,bg=GRID,width=2,height=26).pack(side="left",padx=10)
        Btn(tb,"LOAD FROM EXCEL",self._load,"#5D7C8A","#455A64").pack(side="left",padx=4)
        Btn(tb,"SAVE TO EXCEL",self._save,"#5D7C8A","#455A64").pack(side="left",padx=4)
        Btn(tb,"CLEAR FILTERS",self.grid_w.clear_filters,"#9E9E9E","#757575").pack(side="left",padx=4)
        self.lbl_n=tk.Label(tb,text="0 rows",font=("Segoe UI",10,"bold"),bg=BG,fg=DARK)
        self.lbl_n.pack(side="right")

        # Import bar
        ib=tk.Frame(self,bg=DARK); ib.pack(fill="x",padx=0,pady=(4,0))
        inner2=tk.Frame(ib,bg=DARK); inner2.pack(fill="x",padx=10,pady=8)
        self.btn_imp=Btn(inner2,"IMPORT ALL FROM LCSC",self._start,COPPER,COPPER_D)
        self.btn_imp.pack(side="left")
        self.pbar=ttk.Progressbar(inner2,length=340,mode="determinate")
        self.pbar.pack(side="left",padx=12,fill="x",expand=True)
        self.lbl_p=tk.Label(inner2,text="",font=("Segoe UI",10,"bold"),bg=DARK,fg=TXT_INV)
        self.lbl_p.pack(side="right")

        self.log=LogBox(self,height=6); self.log.pack(fill="x",padx=10,pady=(6,10))

    def _upd(self): self.lbl_n.config(text=f"{self.grid_w.count()} rows")
    def _add_row(self): self.grid_w.add_row(["", "", "Misc", ""]); self._upd()
    def _quick_add(self):
        l=self.v_lcsc.get().strip()
        if not l: return
        self.grid_w.add_row([l,self.v_name.get().strip(),norm_cat(self.v_cat.get()),""])
        self.v_lcsc.set(""); self.v_name.set(""); self._upd()

    def _load(self):
        rows=excel_read_sheet(SHEET_Q,4)
        for r in rows: r[2]=norm_cat(r[2])
        self.grid_w.load(rows); self._upd()
        self.log.log(f"Loaded {len(rows)} rows from Excel","dim")

    def _save(self):
        rows=[r for r in self.grid_w.get_rows() if r[0]]
        if excel_write_sheet(SHEET_Q,rows):
            self.log.log(f"Saved {len(rows)} rows to Excel","ok")
        else: self.log.log("Save failed — check parts.xlsx","err")

    def _start(self):
        if self._running: return
        rows=[r for r in self.grid_w.get_rows() if r[0].strip()]
        if not rows: messagebox.showinfo("Queue empty","Add LCSC parts first."); return
        ez=find_easyeda()
        if not ez:
            messagebox.showerror("Not found","easyeda2kicad not found.\nRun: pip install easyeda2kicad")
            return
        self._running=True
        self.btn_imp.config(state="disabled",text="IMPORTING...")
        self.log.clear(); self.pbar["value"]=0; self.pbar["maximum"]=len(rows)
        mq=TQ.Queue()
        def work():
            ensure_dirs(); ok=[]; fail=[]
            for i,(lcsc,name,cat,note) in enumerate(rows,1):
                mq.put(("log",f"[{i}/{len(rows)}] {lcsc}  {name}","info"))
                sf=sym_path(cat or DEF_CAT); ensure_sym(sf); clear_temp()
                res=subprocess.run([ez,"--full",f"--lcsc_id={lcsc}",
                                    "--output",os.path.join("temp","temp")],
                                   capture_output=True,text=True)
                if res.returncode!=0:
                    mq.put(("log",f"    FAILED: {res.stderr.strip()[:110]}","err"))
                    fail.append([lcsc,name,cat,note]); mq.put(("prog",i,None)); continue
                blocks=extract_blocks(os.path.join(temp_dir(),"temp.kicad_sym"))
                exist={b["name"] for b in extract_blocks(sf)}
                s_ok=False
                for b in blocks:
                    if b["name"] not in exist: merge_block(b["content"],sf); s_ok=True
                fpd=os.path.join(temp_dir(),"temp.pretty"); f_ok=os.path.exists(fpd)
                if f_ok:
                    for fn in os.listdir(fpd):
                        d=os.path.join(fp_dir(),fn)
                        if not os.path.exists(d): shutil.copy2(os.path.join(fpd,fn),d)
                m3d=os.path.join(temp_dir(),"temp.3dshapes"); m_ok=os.path.exists(m3d)
                if m_ok:
                    for fn in os.listdir(m3d):
                        d=os.path.join(m3d_dir(),fn)
                        if not os.path.exists(d): shutil.copy2(os.path.join(m3d,fn),d)
                excel_append_lib(lcsc,name,cat,note,s_ok,f_ok,m_ok)
                ok.append(i-1)
                mq.put(("log",f"    OK   symbol={s_ok}  footprint={f_ok}  3d={m_ok}","ok"))
                mq.put(("prog",i,None))
            clear_temp(); excel_write_sheet(SHEET_Q,fail)
            mq.put(("done",ok,fail))
        threading.Thread(target=work,daemon=True).start(); self._poll(mq)

    def _poll(self,mq):
        try:
            while True:
                m=mq.get_nowait()
                if m[0]=="log": self.log.log(m[1],m[2])
                elif m[0]=="prog":
                    self.pbar["value"]=m[1]
                    self.lbl_p.config(text=f"{m[1]} / {int(self.pbar['maximum'])}")
                elif m[0]=="done":
                    ok,fail=m[1],m[2]
                    self.grid_w.load(fail)
                    self.log.log(f"Finished: {len(ok)} imported, {len(fail)} failed",
                                 "ok" if not fail else "warn")
                    self.btn_imp.config(state="normal",text="IMPORT ALL FROM LCSC")
                    self._running=False; self._upd()
                    self._status(f"Import complete — {len(ok)} OK, {len(fail)} failed")
                    return
        except TQ.Empty: pass
        self.after(100,self._poll,mq)

# ==============================================================================
# CUSTOM SYMBOL TAB
# ==============================================================================
class CustomTab(tk.Frame):
    COLS=["Pin","Name","Type","Side"]; W=[60,190,110,80]
    def __init__(self,parent,status):
        super().__init__(parent,bg=BG); self._status=status; self._build()

    def _build(self):
        top=tk.Frame(self,bg=BG); top.pack(fill="x",padx=10,pady=(10,6))
        tk.Label(top,text="CUSTOM SYMBOL",font=("Segoe UI",13,"bold"),bg=BG,fg=DARK).pack(side="left")
        tk.Label(top,text="Draw a symbol manually — all pins spaced 100 mils apart",
                 font=("Segoe UI",9),bg=BG,fg="#5D5952").pack(side="left",padx=14)

        mf=tk.Frame(self,bg=PANEL,bd=1,relief="solid"); mf.pack(fill="x",padx=10,pady=(0,8))
        r1=tk.Frame(mf,bg=PANEL); r1.pack(fill="x",padx=10,pady=(8,4))
        tk.Label(r1,text="Symbol Name:",font=("Segoe UI",9,"bold"),bg=PANEL,fg=TXT).pack(side="left")
        self.v_name=tk.StringVar()
        tk.Entry(r1,textvariable=self.v_name,width=22,font=("Segoe UI",10),
                 relief="solid",bd=1).pack(side="left",padx=(4,14))
        tk.Label(r1,text="Category:",font=("Segoe UI",9,"bold"),bg=PANEL,fg=TXT).pack(side="left")
        self.v_cat=tk.StringVar(value="ICs")
        ttk.Combobox(r1,textvariable=self.v_cat,values=CATS,width=14,
                     state="readonly").pack(side="left",padx=(4,14))
        tk.Label(r1,text="Reference:",font=("Segoe UI",9,"bold"),bg=PANEL,fg=TXT).pack(side="left")
        self.v_ref=tk.StringVar(value="U")
        ttk.Combobox(r1,textvariable=self.v_ref,values=REFS,width=6).pack(side="left",padx=4)

        r2=tk.Frame(mf,bg=PANEL); r2.pack(fill="x",padx=10,pady=(0,8))
        tk.Label(r2,text="Description:",font=("Segoe UI",9,"bold"),bg=PANEL,fg=TXT).pack(side="left")
        self.v_desc=tk.StringVar()
        tk.Entry(r2,textvariable=self.v_desc,font=("Segoe UI",10),
                 relief="solid",bd=1).pack(side="left",fill="x",expand=True,padx=(4,0))

        tb=tk.Frame(self,bg=BG); tb.pack(fill="x",padx=10,pady=(0,4))
        Btn(tb,"+ PIN",self._add_pin,SLATE,SLATE_D).pack(side="left",padx=(0,4))
        Btn(tb,"DELETE",lambda:self.grid_w.delete_selected(),"#8D6E63","#6D4C41").pack(side="left",padx=4)
        Btn(tb,"CLEAR",self.grid_w_clear,"#78909C","#546E7A").pack(side="left",padx=4)
        Btn(tb,"PASTE",lambda:self.grid_w.paste_clipboard(),SLATE,SLATE_D).pack(side="left",padx=4)
        tk.Label(tb,text="Paste columns: Pin | Name | Type | Side",
                 font=("Segoe UI",9,"bold"),bg=BG,fg=COPPER).pack(side="right")

        self.grid_w=DataGrid(self,self.COLS,self.W,
                             combos={"Type":PTYPES,"Side":PSIDES},filters=False)
        self.grid_w.pack(fill="both",expand=True,padx=10)
        for i in range(1,5): self.grid_w.add_row([str(i),"","IN","L"])

        af=tk.Frame(self,bg=DARK); af.pack(fill="x",pady=(6,0))
        inner=tk.Frame(af,bg=DARK); inner.pack(fill="x",padx=10,pady=8)
        Btn(inner,"CREATE SYMBOL",self._create,MOSS,MOSS_D).pack(side="left")
        self.lbl_res=tk.Label(inner,text="",font=("Segoe UI",10,"bold"),bg=DARK,fg=TXT_INV)
        self.lbl_res.pack(side="left",padx=14)

        self.log=LogBox(self,height=5); self.log.pack(fill="x",padx=10,pady=(6,10))

    def grid_w_clear(self): self.grid_w.clear()
    def _add_pin(self): self.grid_w.add_row([str(self.grid_w.count()+1),"","IN","L"])

    def _create(self):
        name=self.v_name.get().strip(); cat=self.v_cat.get().strip()
        ref=self.v_ref.get().strip() or "U"; desc=self.v_desc.get().strip()
        if not name: messagebox.showwarning("Missing","Enter a symbol name."); return
        pins=[]
        for num,pn,pt,sd in self.grid_w.get_rows():
            if num and pn:
                pins.append({"number":num,"name":pn,"type":pt.lower(),
                             "side":sd.upper()[:1] if sd else "L"})
        if not pins: messagebox.showwarning("Missing","Add at least one pin."); return
        ensure_dirs(); sf=sym_path(cat); ensure_sym(sf)
        if sym_exists(sf,name):
            if not messagebox.askyesno("Exists",f"'{name}' already exists. Overwrite?"): return
            with open(sf,"r",encoding="utf-8") as f: c=f.read()
            s=c.find(f'(symbol "{name}"')
            if s!=-1:
                st,e=0,-1
                for i in range(s,len(c)):
                    if c[i]=="(": st+=1
                    elif c[i]==")":
                        st-=1
                        if st==0: e=i+1; break
                if e!=-1:
                    c=c[:s]+c[e:]
                    with open(sf,"w",encoding="utf-8") as f: f.write(c)
        try:
            merge_block(gen_kicad_sym(name,ref,desc,pins),sf)
            excel_append_lib(name,name,cat,desc,True,False,False)
            self.lbl_res.config(text=f"Created: {name}",fg="#AED581")
            L=sum(1 for p in pins if p["side"]=="L"); R=sum(1 for p in pins if p["side"]=="R")
            T=sum(1 for p in pins if p["side"]=="T"); B=sum(1 for p in pins if p["side"]=="B")
            self.log.log(f"Symbol '{name}' -> {sf}","ok")
            self.log.log(f"    {len(pins)} pins   Left={L} Right={R} Top={T} Bottom={B}","info")
            self._status(f"Symbol '{name}' created")
        except Exception as ex:
            self.lbl_res.config(text=f"Error: {ex}",fg="#EF9A9A")
            self.log.log(f"ERROR: {ex}","err")

# ==============================================================================
# LIBRARY TAB
# ==============================================================================
class LibTab(tk.Frame):
    COLS=["LCSC","Name","Category","Note","Sym","FP","3D","Updated"]
    W=[100,180,120,230,45,45,45,120]
    def __init__(self,parent,status):
        super().__init__(parent,bg=BG); self._status=status; self._build(); self._load()

    def _build(self):
        top=tk.Frame(self,bg=BG); top.pack(fill="x",padx=10,pady=(10,6))
        tk.Label(top,text="LIBRARY",font=("Segoe UI",13,"bold"),bg=BG,fg=DARK).pack(side="left")
        tk.Label(top,text="Click a row number on the left to select, then use the buttons above",
                 font=("Segoe UI",9),bg=BG,fg="#5D5952").pack(side="left",padx=14)
        self.lbl_n=tk.Label(top,text="",font=("Segoe UI",10,"bold"),bg=BG,fg=DARK)
        self.lbl_n.pack(side="right")

        # Compact toolbar: one dropdown for folders, real per-row actions
        fb=tk.Frame(self,bg=PANEL,bd=1,relief="solid"); fb.pack(fill="x",padx=10,pady=(0,6))
        inner=tk.Frame(fb,bg=PANEL); inner.pack(fill="x",padx=10,pady=7)

        mb=tk.Menubutton(inner,text="OPEN FOLDER  \u25BE",bg="#6D8B74",fg=TXT_INV,
                         font=("Segoe UI",9,"bold"),relief="flat",bd=0,
                         padx=12,pady=6,cursor="hand2",activebackground="#55705C")
        menu=tk.Menu(mb,tearoff=0,font=("Segoe UI",9),bg=CELL_BG,fg=TXT,
                     activebackground=COPPER,activeforeground=TXT_INV)
        menu.add_command(label="Library root",command=lambda:open_path(LIB_ROOT))
        menu.add_separator()
        menu.add_command(label="symbols",     command=lambda:open_path(sym_dir(),create=True))
        menu.add_command(label="footprints",  command=lambda:open_path(fp_dir(),create=True))
        menu.add_command(label="3d models",   command=lambda:open_path(m3d_dir(),create=True))
        mb.config(menu=menu); mb.pack(side="left")

        tk.Frame(inner,bg=GRID,width=2,height=26).pack(side="left",padx=12)
        tk.Label(inner,text="Selected component:",font=("Segoe UI",9,"bold"),
                 bg=PANEL,fg=TXT).pack(side="left",padx=(0,8))
        Btn(inner,"REVEAL FILES",self._reveal_files,COPPER,COPPER_D).pack(side="left",padx=3)
        Btn(inner,"EDIT IN KICAD",self._open_sym_file,COPPER,COPPER_D).pack(side="left",padx=3)
        Btn(inner,"DELETE COMPONENT",self._delete_component,"#A03A2E","#7E2C22").pack(side="left",padx=3)

        self.lbl_paths=tk.Label(self,text="",font=("Segoe UI",8),bg=BG,fg="#6B665E",anchor="w")
        self.lbl_paths.pack(fill="x",padx=10,pady=(0,4))
        self._show_paths()

        self.grid_w=DataGrid(self,self.COLS,self.W,
                             combos={"Category":CATS,"Sym":["v","x"],
                                     "FP":["v","x"],"3D":["v","x"]},
                             on_change=self._upd)
        self.grid_w.pack(fill="both",expand=True,padx=10)

        tb=tk.Frame(self,bg=BG); tb.pack(fill="x",padx=10,pady=8)
        Btn(tb,"RELOAD",self._load,SLATE,SLATE_D).pack(side="left",padx=(0,4))
        Btn(tb,"SAVE TO EXCEL",self._save,MOSS,MOSS_D).pack(side="left",padx=4)
        Btn(tb,"DELETE ROW",self.grid_w.delete_selected,"#8D6E63","#6D4C41").pack(side="left",padx=4)
        Btn(tb,"CLEAR FILTERS",self.grid_w.clear_filters,"#9E9E9E","#757575").pack(side="left",padx=4)
        Btn(tb,"RESCAN LIBRARY",self._rescan,"#5D7C8A","#455A64").pack(side="left",padx=4)
        Btn(tb,"SELECT ALL",self.grid_w.select_all,"#9E9E9E","#757575").pack(side="left",padx=4)
        self.lbl_msg=tk.Label(tb,text="",font=("Segoe UI",9,"bold"),bg=BG,fg=MOSS_D)
        self.lbl_msg.pack(side="left",padx=14)

    def _upd(self): self.lbl_n.config(text=f"{self.grid_w.count()} components")

    def _show_paths(self):
        def mark(p): return ("OK  " if os.path.isdir(p) else "MISSING  ")+p
        self.lbl_paths.config(text=(f"symbols: {mark(sym_dir())}     "
                                    f"footprints: {mark(fp_dir())}     "
                                    f"3d: {mark(m3d_dir())}"))

    def _load(self):
        rows=excel_read_sheet(SHEET_LIB,8)
        self.grid_w.load(rows); self._upd(); self._show_paths()
        self.lbl_msg.config(text=f"Loaded {len(rows)} rows",fg=MOSS_D)

    def _save(self):
        rows=[r for r in self.grid_w.get_rows() if r[0]]
        if excel_write_sheet(SHEET_LIB,rows,has_status=True):
            self.lbl_msg.config(text=f"Saved {len(rows)} rows to Excel",fg=MOSS_D)
            self._status(f"Library saved — {len(rows)} components")
        else:
            self.lbl_msg.config(text="Save failed — check parts.xlsx",fg=ERR)

    def _sel_row(self):
        rows=self.grid_w.selected_rows()
        if not rows:
            messagebox.showinfo("No row selected",
                "Click the number on the left of a row to select it first.")
            return None
        return rows[0]

    def _search_dir(self,directory,keys,exts):
        if not os.path.exists(directory): return None
        files=[f for f in os.listdir(directory) if f.lower().endswith(exts)]
        for k in keys:
            if not k: continue
            k2=k.lower().replace(" ","").replace("-","").replace("_","")
            for f in files:
                f2=f.lower().replace(" ","").replace("-","").replace("_","")
                if k2 and k2 in f2: return os.path.join(directory,f)
        return None

    def _reveal_files(self):
        """Show every file linked to the selected component."""
        r=self._sel_row()
        if not r: return
        lcsc,name,cat=r[0],r[1],(r[2] or DEF_CAT)
        info=resolve_component(name,lcsc,cat)

        dlg=tk.Toplevel(self); dlg.title(f"Files — {name or lcsc}")
        dlg.configure(bg=BG); dlg.resizable(False,False); dlg.grab_set()
        tk.Label(dlg,text=name or lcsc,font=("Segoe UI",12,"bold"),
                 bg=BG,fg=DARK).pack(anchor="w",padx=16,pady=(14,2))
        tk.Label(dlg,text=f"{lcsc}    {cat}",font=("Segoe UI",9),
                 bg=BG,fg="#5D5952").pack(anchor="w",padx=16,pady=(0,10))

        rows=[
            ("Symbol library",
             os.path.basename(info["sym_file"]) if info["sym_file"] else None,
             info["sym_file"], sym_dir()),
            ("Footprint",
             (os.path.basename(info["fp_path"]) if info["fp_path"]
              else (info["fp_name"]+".kicad_mod  (file missing)" if info["fp_name"] else None)),
             info["fp_path"], fp_dir()),
            ("3D model",
             (os.path.basename(info["model_path"]) if info["model_path"]
              else (info["model_name"]+"  (file missing)" if info["model_name"] else None)),
             info["model_path"], m3d_dir()),
        ]
        for label,shown,path,folder in rows:
            row=tk.Frame(dlg,bg=PANEL,bd=1,relief="solid")
            row.pack(fill="x",padx=16,pady=3)
            tk.Label(row,text=label,font=("Segoe UI",9,"bold"),width=14,
                     anchor="w",bg=PANEL,fg=TXT).pack(side="left",padx=8,pady=8)
            if path:
                tk.Label(row,text=shown,font=("Segoe UI",9),bg=PANEL,fg=MOSS_D,
                         width=42,anchor="w").pack(side="left")
                Btn(row,"SHOW",lambda q=path:reveal_file(q),SLATE,SLATE_D).pack(side="right",padx=8)
            else:
                tk.Label(row,text=shown or "not linked",font=("Segoe UI",9,"bold"),
                         bg=PANEL,fg=WARN,width=42,anchor="w").pack(side="left")
                Btn(row,"OPEN FOLDER",lambda f=folder:open_path(f,create=True),
                    COPPER,COPPER_D).pack(side="right",padx=8)

        if info["sym_file"] and not info["fp_name"]:
            tk.Label(dlg,text="This symbol has no Footprint field — assign one in KiCad "
                             "Symbol Editor so the link can be followed.",
                     font=("Segoe UI",8),bg=BG,fg=WARN,wraplength=420,
                     justify="left").pack(anchor="w",padx=16,pady=(8,0))
        Btn(dlg,"CLOSE",dlg.destroy,SLATE,SLATE_D).pack(pady=14)

    def _open_sym_file(self):
        r=self._sel_row()
        if not r: return
        sf=sym_path(r[2] or DEF_CAT)
        if os.path.exists(sf):
            open_path(sf)
            self.lbl_msg.config(text=f"Opened {os.path.basename(sf)}",fg=MOSS_D)
        else:
            open_path(sym_dir(),create=True)
            self.lbl_msg.config(text=f"{os.path.basename(sf)} not found",fg=WARN)

    def _delete_component(self):
        """Remove the symbol from its .kicad_sym, delete its files, drop the row."""
        idx=self.grid_w.selected_index()
        r=self._sel_row()
        if r is None or idx is None: return
        lcsc,name,cat=r[0],r[1],(r[2] or DEF_CAT)
        info=resolve_component(name,lcsc,cat)

        detail=[]
        if info["sym_file"]: detail.append(f"Symbol: {info['sym_name']}  in  {os.path.basename(info['sym_file'])}")
        if info["fp_path"]:  detail.append(f"Footprint: {os.path.basename(info['fp_path'])}")
        if info["model_path"]: detail.append(f"3D model: {os.path.basename(info['model_path'])}")
        if not detail: detail.append("No files found — only the row will be removed.")

        if not messagebox.askyesno("Delete component",
            f"Permanently delete '{name or lcsc}'?\n\n"+"\n".join(detail)+
            "\n\nThis cannot be undone."):
            return

        removed=[]
        if info["sym_file"] and info["sym_name"]:
            try:
                c=open(info["sym_file"],encoding="utf-8").read()
                st=c.find(f'(symbol "{info["sym_name"]}"')
                if st!=-1:
                    d,e=0,-1
                    for i in range(st,len(c)):
                        if c[i]=="(": d+=1
                        elif c[i]==")":
                            d-=1
                            if d==0: e=i+1; break
                    if e!=-1:
                        open(info["sym_file"],"w",encoding="utf-8").write(c[:st]+c[e:])
                        removed.append("symbol")
            except Exception as ex:
                self.lbl_msg.config(text=f"Symbol removal failed: {ex}",fg=ERR)

        if messagebox.askyesno("Shared files",
            "Also delete the footprint and 3D model files?\n\n"
            "Choose No if other components may use the same package."):
            for pth,tag in ((info["fp_path"],"footprint"),(info["model_path"],"3d model")):
                if pth and os.path.exists(pth):
                    try: os.remove(pth); removed.append(tag)
                    except Exception: pass

        self.grid_w.delete_selected(); self._save()
        self.lbl_msg.config(text=f"Deleted {name or lcsc}  ({', '.join(removed) or 'row only'})",
                            fg=MOSS_D)

    def _rescan(self):
        """Re-check symbol, footprint and 3D presence for every row via real links."""
        if not os.path.isdir(sym_dir()):
            self.lbl_msg.config(text="symbols folder not found",fg=ERR); return
        n_s=n_f=n_m=0
        for row in self.grid_w._rows:
            lcsc=row["vars"][0].get().strip()
            name=row["vars"][1].get().strip()
            cat =row["vars"][2].get().strip() or DEF_CAT
            info=resolve_component(name,lcsc,cat)
            row["vars"][4].set("v" if info["sym_file"] else "x")
            row["vars"][5].set("v" if info["fp_path"] else "x")
            row["vars"][6].set("v" if info["model_path"] else "x")
            n_s+=1 if info["sym_file"] else 0
            n_f+=1 if info["fp_path"] else 0
            n_m+=1 if info["model_path"] else 0
        self.lbl_msg.config(
            text=f"Rescanned — symbols {n_s}, footprints {n_f}, 3D {n_m}. Click Save to keep.",
            fg=MOSS_D)
        self._status(f"Rescan: {n_s} symbols, {n_f} footprints, {n_m} models")


# ==============================================================================
# TOOLS TAB — library maintenance
# ==============================================================================
def scan_library():
    """Collect every symbol, footprint and 3D file plus any problems found."""
    rep={"symbols":{}, "footprints":[], "models":[], "issues":[],
         "dupes":[], "by_cat":{}}
    if os.path.isdir(sym_dir()):
        for fn in sorted(os.listdir(sym_dir())):
            if not (fn.startswith(f"{LIB_PFX}-") and fn.endswith(".kicad_sym")): continue
            cat=fn[len(LIB_PFX)+1:-len(".kicad_sym")]
            blocks=extract_blocks(os.path.join(sym_dir(),fn))
            rep["by_cat"][cat]=len(blocks)
            for b in blocks:
                nm=b["name"]
                if nm in rep["symbols"]:
                    rep["dupes"].append(f"{nm}  (in {rep['symbols'][nm]['cat']} and {cat})")
                fpm=re.search(r'\(property\s+"Footprint"\s+"([^"]*)"',b["content"])
                rep["symbols"][nm]={"cat":cat,"lcsc":b["lcsc"],
                                    "fp":fpm.group(1) if fpm else ""}
    if os.path.isdir(fp_dir()):
        rep["footprints"]=[f for f in sorted(os.listdir(fp_dir()))
                           if f.endswith(".kicad_mod")]
    if os.path.isdir(m3d_dir()):
        rep["models"]=[f for f in sorted(os.listdir(m3d_dir()))
                       if f.lower().endswith((".step",".stp",".wrl"))]

    fp_names={os.path.splitext(f)[0].lower() for f in rep["footprints"]}
    # symbols with no footprint assigned, or pointing at a file that is not there
    for nm,info in rep["symbols"].items():
        fp=info["fp"]
        if not fp:
            rep["issues"].append(("symbol has no Footprint field",nm))
        else:
            base=fp.split(":")[-1].lower()
            if base and base not in fp_names:
                rep["issues"].append(("footprint file missing",f"{nm}  ->  {fp}"))
    # footprints whose 3D model path points nowhere
    model_names={f.lower() for f in rep["models"]}
    for f in rep["footprints"]:
        try: c=open(os.path.join(fp_dir(),f),encoding="utf-8").read()
        except Exception: continue
        for m in re.finditer(r'\(model\s+"([^"]+)"',c):
            path=m.group(1)
            base=os.path.basename(path).lower()
            if base not in model_names:
                rep["issues"].append(("3D model missing",f"{f}  ->  {base}"))
            if "${" not in path:
                rep["issues"].append(("3D path not using a variable",f))
    return rep


class ToolsTab(tk.Frame):
    def __init__(self,parent,status):
        super().__init__(parent,bg=BG); self._status=status; self._build()

    def _build(self):
        top=tk.Frame(self,bg=BG); top.pack(fill="x",padx=10,pady=(10,6))
        tk.Label(top,text="TOOLS",font=("Segoe UI",13,"bold"),bg=BG,fg=DARK).pack(side="left")
        tk.Label(top,text="Check and repair the library, then export what you need",
                 font=("Segoe UI",9),bg=BG,fg="#5D5952").pack(side="left",padx=14)

        cards=tk.Frame(self,bg=BG); cards.pack(fill="x",padx=10)
        specs=[
            ("HEALTH CHECK","Find symbols with no footprint, missing files and broken 3D paths",
             self._health,COPPER,COPPER_D),
            ("FIX 3D PATHS","Rewrite every model path to ${KICAD_USER_LIB} so both machines resolve it",
             self._fix3d,"#6D8B74","#55705C"),
            ("FIND DUPLICATES","List symbols or LCSC codes that appear in more than one file",
             self._dupes,"#7C6A9C","#61527C"),
            ("STATISTICS","Count symbols per category, footprints and 3D models",
             self._stats,SLATE,SLATE_D),
            ("EXPORT CSV","Write the whole library to a CSV file for BOM or procurement",
             self._csv,"#5D7C8A","#455A64"),
            ("KICAD LIB TABLE","Generate sym-lib-table and fp-lib-table entries to paste into KiCad",
             self._libtable,"#8D6E63","#6D4C41"),
        ]
        for i,(title,desc,cmd,c1,c2) in enumerate(specs):
            row,col=divmod(i,3)
            card=tk.Frame(cards,bg=PANEL,bd=1,relief="solid")
            card.grid(row=row,column=col,sticky="nsew",padx=4,pady=4)
            cards.grid_columnconfigure(col,weight=1)
            Btn(card,title,cmd,c1,c2).pack(fill="x",padx=10,pady=(10,6))
            tk.Label(card,text=desc,font=("Segoe UI",8),bg=PANEL,fg="#5D5952",
                     wraplength=290,justify="left").pack(anchor="w",padx=10,pady=(0,10))

        tk.Label(self,text="Output",font=("Segoe UI",9,"bold"),
                 bg=BG,fg=DARK).pack(anchor="w",padx=10,pady=(10,2))
        self.log=LogBox(self,height=16); self.log.pack(fill="both",expand=True,padx=10,pady=(0,10))

    # ── Actions ───────────────────────────────────────────────────────────────
    def _health(self):
        self.log.clear(); self.log.log("HEALTH CHECK","info")
        self.log.log("-"*60,"dim")
        r=scan_library()
        self.log.log(f"Symbols    {len(r['symbols'])}","ok")
        self.log.log(f"Footprints {len(r['footprints'])}","ok")
        self.log.log(f"3D models  {len(r['models'])}","ok")
        self.log.log("","dim")
        if not r["issues"]:
            self.log.log("No problems found.","ok"); return
        groups={}
        for kind,detail in r["issues"]: groups.setdefault(kind,[]).append(detail)
        for kind,items in groups.items():
            self.log.log(f"{kind}  ({len(items)})","warn")
            for it in items[:40]: self.log.log(f"    {it}","dim")
            if len(items)>40: self.log.log(f"    ... and {len(items)-40} more","dim")
            self.log.log("","dim")
        self._status(f"Health check: {len(r['issues'])} issues")

    def _fix3d(self):
        if not os.path.isdir(fp_dir()):
            self.log.log("footprints folder not found","err"); return
        if not messagebox.askyesno("Fix 3D paths",
            "Rewrite the 3D model path inside every .kicad_mod so it uses "
            "${KICAD_USER_LIB}?\n\nThe file name is kept, only the folder part changes."):
            return
        self.log.clear(); self.log.log("FIX 3D PATHS","info"); self.log.log("-"*60,"dim")
        pat=re.compile(r'\(model\s+"[^"]*?([^"/\\]+\.(?:step|stp|wrl))"',re.I)
        target=r'(model "${KICAD_USER_LIB}/3dmodels/'+LIB_PFX+r'.3dshapes/\1"'
        n=0
        for f in sorted(os.listdir(fp_dir())):
            if not f.endswith(".kicad_mod"): continue
            fp=os.path.join(fp_dir(),f)
            try:
                c=open(fp,encoding="utf-8").read()
                new=pat.sub(target,c)
                if new!=c:
                    open(fp,"w",encoding="utf-8").write(new)
                    self.log.log(f"  fixed  {f}","ok"); n+=1
            except Exception as ex:
                self.log.log(f"  error  {f}: {ex}","err")
        self.log.log("","dim")
        self.log.log(f"Updated {n} footprint files.","ok" if n else "dim")
        self._status(f"Fixed 3D paths in {n} files")

    def _dupes(self):
        self.log.clear(); self.log.log("DUPLICATES","info"); self.log.log("-"*60,"dim")
        r=scan_library()
        lcsc_map={}
        for nm,info in r["symbols"].items():
            if info["lcsc"]: lcsc_map.setdefault(info["lcsc"],[]).append(nm)
        dl=[(k,v) for k,v in lcsc_map.items() if len(v)>1]
        if not r["dupes"] and not dl:
            self.log.log("No duplicates found.","ok"); return
        if r["dupes"]:
            self.log.log(f"Duplicate symbol names  ({len(r['dupes'])})","warn")
            for d in r["dupes"]: self.log.log(f"    {d}","dim")
        if dl:
            self.log.log("","dim")
            self.log.log(f"Same LCSC used by several symbols  ({len(dl)})","warn")
            for k,v in dl: self.log.log(f"    {k}  ->  {', '.join(v)}","dim")
        self._status(f"{len(r['dupes'])+len(dl)} duplicates")

    def _stats(self):
        self.log.clear(); self.log.log("STATISTICS","info"); self.log.log("-"*60,"dim")
        r=scan_library()
        if not r["by_cat"]:
            self.log.log("No symbol files found under symbols/","warn"); return
        total=0
        for cat,n in sorted(r["by_cat"].items(),key=lambda x:-x[1]):
            bar="#"*min(n,40)
            self.log.log(f"  {cat:<16}{n:>4}  {bar}","ok" if n else "dim"); total+=n
        self.log.log("","dim")
        self.log.log(f"  {'TOTAL':<16}{total:>4}","info")
        self.log.log(f"  {'footprints':<16}{len(r['footprints']):>4}","info")
        self.log.log(f"  {'3d models':<16}{len(r['models']):>4}","info")
        cov=(len(r['footprints'])/total*100) if total else 0
        self.log.log("","dim")
        self.log.log(f"  Footprint coverage  {cov:.0f}%","ok" if cov>80 else "warn")

    def _csv(self):
        rows=excel_read_sheet(SHEET_LIB,8)
        if not rows:
            self.log.log("Library sheet is empty","warn"); return
        path=filedialog.asksaveasfilename(defaultextension=".csv",
                initialfile="library_export.csv",
                filetypes=[("CSV file","*.csv")], initialdir=LIB_ROOT)
        if not path: return
        try:
            import csv as _csv
            with open(path,"w",newline="",encoding="utf-8-sig") as f:
                w=_csv.writer(f)
                w.writerow(["LCSC","Name","Category","Description",
                            "Symbol","Footprint","3D","Updated"])
                w.writerows(rows)
            self.log.clear()
            self.log.log(f"Exported {len(rows)} components","ok")
            self.log.log(path,"dim")
            self._status(f"Exported {len(rows)} rows to CSV")
            if messagebox.askyesno("Export complete","Open the file now?"): open_path(path)
        except Exception as ex:
            self.log.log(f"Export failed: {ex}","err")

    def _libtable(self):
        self.log.clear(); self.log.log("KICAD LIBRARY TABLE","info"); self.log.log("-"*60,"dim")
        syms=[f for f in sorted(os.listdir(sym_dir()))
              if f.endswith(".kicad_sym")] if os.path.isdir(sym_dir()) else []
        if not syms:
            self.log.log("No .kicad_sym files found under symbols/","warn"); return
        lines=["(sym_lib_table","  (version 7)"]
        for f in syms:
            nick=f[:-len(".kicad_sym")]
            lines.append(f'  (lib (name "{nick}")(type "KiCad")'
                         f'(uri "${{KICAD_USER_LIB}}/symbols/{f}")(options "")(descr ""))')
        lines.append(")")
        lines.append("")
        lines.append("(fp_lib_table")
        lines.append("  (version 7)")
        pretty=os.path.basename(fp_dir())
        rel="footprints/"+pretty if pretty.endswith(".pretty") else "footprints"
        lines.append(f'  (lib (name "{LIB_PFX}")(type "KiCad")'
                     f'(uri "${{KICAD_USER_LIB}}/{rel}")(options "")(descr ""))')
        lines.append(")")
        text="\n".join(lines)
        for l in lines: self.log.log(l,"dim")
        self.log.log("","dim")
        self.log.log("Set KICAD_USER_LIB in KiCad: Preferences > Configure Paths","info")
        self.log.log(f"    KICAD_USER_LIB  =  {LIB_ROOT}","ok")
        try:
            self.clipboard_clear(); self.clipboard_append(text)
            self.log.log("","dim"); self.log.log("Copied to clipboard.","ok")
        except Exception: pass
        self._status("Library table generated")


# ==============================================================================
# APP
# ==============================================================================
class App:
    def __init__(self):
        self.root=tk.Tk()
        self.root.title(f"{APP_TITLE}  {APP_VER}")
        self.root.geometry("1080x760"); self.root.minsize(900,640)
        self.root.configure(bg=BG)
        apply_icon(self.root)
        self._style(); self._build()

    def _style(self):
        s=ttk.Style()
        try: s.theme_use("clam")
        except Exception: pass
        s.configure("TNotebook",background=DARK,borderwidth=0)
        s.configure("TNotebook.Tab",font=("Segoe UI",10,"bold"),padding=(20,9),
                    background="#455A64",foreground="#CFD8DC",borderwidth=0)
        s.map("TNotebook.Tab",background=[("selected",BG)],
              foreground=[("selected",DARK)])
        s.configure("TCombobox",fieldbackground=CELL_BG,background=CELL_BG)
        s.configure("Horizontal.TProgressbar",background=COPPER,
                    troughcolor="#1F2C33",borderwidth=0,thickness=18)

    def _build(self):
        hb=tk.Frame(self.root,bg=DARK,height=56); hb.pack(fill="x"); hb.pack_propagate(False)

        # Logo, if app_logo.png sits next to the program
        logo=resource("app_logo.png")
        if logo:
            try:
                img=tk.PhotoImage(file=logo)
                f=max(1,img.width()//40)
                img=img.subsample(f,f)
                lbl=tk.Label(hb,image=img,bg=DARK); lbl.image=img
                lbl.pack(side="left",padx=(14,8))
            except Exception: pass

        tk.Label(hb,text=APP_TITLE,font=("Segoe UI",15,"bold"),
                 bg=DARK,fg=TXT_INV).pack(side="left",padx=(2,0))
        tk.Label(hb,text=APP_VER,font=("Segoe UI",9,"bold"),
                 bg=DARK,fg=COPPER).pack(side="left",padx=(6,0))
        tk.Button(hb,text="?",command=self._about,font=("Segoe UI",9,"bold"),
                  bg=DARK,fg="#78909C",relief="flat",bd=0,cursor="hand2",
                  activebackground=DARK,activeforeground=COPPER,
                  padx=8).pack(side="left",padx=(10,0))
        ff=tk.Frame(hb,bg=DARK); ff.pack(side="right",padx=16)
        tk.Label(ff,text="Library folder:",font=("Segoe UI",9,"bold"),
                 bg=DARK,fg="#B0BEC5").pack(side="left")
        self.v_root=tk.StringVar(value=LIB_ROOT)
        e_root=tk.Entry(ff,textvariable=self.v_root,width=38,font=("Segoe UI",9),
                        relief="flat",bd=2)
        e_root.pack(side="left",padx=6)
        e_root.bind("<Return>",lambda ev:self._apply_typed_root())
        Btn(ff,"SET",self._apply_typed_root,MOSS,MOSS_D).pack(side="left",padx=(0,4))
        Btn(ff,"CHANGE",self._browse,SLATE,SLATE_D).pack(side="left")

        self.nb=ttk.Notebook(self.root); self.nb.pack(fill="both",expand=True)
        self.t_q=QueueTab(self.nb,self._status)
        self.t_c=CustomTab(self.nb,self._status)
        self.t_l=LibTab(self.nb,self._status)
        self.t_t=ToolsTab(self.nb,self._status)
        self.nb.add(self.t_q,text="QUEUE IMPORT")
        self.nb.add(self.t_c,text="CUSTOM SYMBOL")
        self.nb.add(self.t_l,text="LIBRARY")
        self.nb.add(self.t_t,text="TOOLS")

        if not looks_like_root(LIB_ROOT):
            warn=tk.Frame(self.root,bg="#8D3B1E",height=30); warn.pack(fill="x")
            warn.pack_propagate(False)
            tk.Label(warn,text="Library folder not detected — click CHANGE at the top right "
                               "and select the folder that contains symbols/ and parts.xlsx",
                     font=("Segoe UI",9,"bold"),bg="#8D3B1E",fg="#FFE0B2").pack(pady=6)

        sb=tk.Frame(self.root,bg=DARK2,height=26); sb.pack(fill="x",side="bottom")
        sb.pack_propagate(False)
        self.lbl_s=tk.Label(sb,text="Ready",font=("Segoe UI",9),bg=DARK2,fg="#B0BEC5")
        self.lbl_s.pack(side="left",padx=12)
        tk.Label(sb,text=f"by {APP_AUTHOR}",font=("Segoe UI",9,"bold"),
                 bg=DARK2,fg=COPPER).pack(side="left",padx=(0,12))
        tk.Label(sb,text="Click row number to select    Ctrl+click add    Shift+click range    "
                         "Ctrl+V paste    Tab / Enter move",
                 font=("Segoe UI",9),bg=DARK2,fg="#78909C").pack(side="right",padx=12)

    def _apply_typed_root(self):
        """Accept a path typed or pasted into the root box."""
        d=self.v_root.get().strip().strip('"')
        if not os.path.isdir(d):
            messagebox.showerror("Not a folder","This path does not exist:\n\n"+d); return
        set_root(d); self.v_root.set(LIB_ROOT)
        self._check_layout(); self.t_q._load(); self.t_l._load()
        self._status(f"Library root: {LIB_ROOT}")

    def _browse(self):
        d=filedialog.askdirectory(title="Select KiCad library folder "
                                        "(the folder containing parts.xlsx and symbols/)",
                                  initialdir=LIB_ROOT)
        if not d: return
        set_root(d); self.v_root.set(LIB_ROOT)
        self._check_layout()
        self.t_q._load(); self.t_l._load()
        self._status(f"Library root: {LIB_ROOT}")

    def _check_layout(self):
        """Warn about anything missing under the chosen root."""
        missing=[n for n,pth in [("parts.xlsx",parts_file()),
                                 ("symbols/",sym_dir()),
                                 ("footprints/",fp_dir()),
                                 ("3dmodels/",m3d_dir())] if not os.path.exists(pth)]
        if missing:
            if messagebox.askyesno("Folders missing",
                "These are missing under:\n"+LIB_ROOT+"\n\n  "+"\n  ".join(missing)+
                "\n\nCreate the missing folders now?"):
                ensure_dirs()
                self._status("Created missing folders")

    def _about(self):
        dlg=tk.Toplevel(self.root); dlg.title("About")
        dlg.configure(bg=BG); dlg.resizable(False,False); dlg.grab_set()
        apply_icon(dlg)

        logo=resource("app_logo.png")
        if logo:
            try:
                img=tk.PhotoImage(file=logo)
                f=max(1,img.width()//96)
                img=img.subsample(f,f)
                l=tk.Label(dlg,image=img,bg=BG); l.image=img; l.pack(pady=(20,10))
            except Exception: pass

        tk.Label(dlg,text=APP_TITLE,font=("Segoe UI",15,"bold"),
                 bg=BG,fg=DARK).pack()
        tk.Label(dlg,text=f"Version {APP_VER}",font=("Segoe UI",10),
                 bg=BG,fg=COPPER).pack(pady=(2,14))

        info=tk.Frame(dlg,bg=PANEL,bd=1,relief="solid"); info.pack(fill="x",padx=24)
        for k,v in [("Author",APP_AUTHOR),
                    ("Year",APP_YEAR),
                    ("Library root",LIB_ROOT)]:
            r=tk.Frame(info,bg=PANEL); r.pack(fill="x",padx=12,pady=5)
            tk.Label(r,text=k,font=("Segoe UI",9,"bold"),width=12,anchor="w",
                     bg=PANEL,fg=TXT).pack(side="left")
            tk.Label(r,text=v,font=("Segoe UI",9),bg=PANEL,fg="#5D5952",
                     anchor="w",wraplength=300,justify="left").pack(side="left")

        tk.Label(dlg,text="Manages KiCad symbol, footprint and 3D model libraries.\n"
                          "Imports parts from LCSC and keeps everything in one place.",
                 font=("Segoe UI",9),bg=BG,fg="#5D5952",justify="center").pack(pady=14)
        Btn(dlg,"CLOSE",dlg.destroy,SLATE,SLATE_D).pack(pady=(0,20))

    def _status(self,m): self.lbl_s.config(text=m)
    def run(self): self.root.mainloop()

if __name__=="__main__":
    base=os.path.dirname(sys.executable) if getattr(sys,"frozen",False) \
         else os.path.dirname(os.path.abspath(__file__))
    os.chdir(base)
    LIB_ROOT=base
    load_root()          # saved root, else auto-detect by walking up
    App().run()
