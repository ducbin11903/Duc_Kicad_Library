import os
import sys
import subprocess
import shutil
import re
import platform
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def check_file_not_open(filepath):
    """
    Kiem tra file Excel co dang bi mo khong.
    """
    if not os.path.exists(filepath):
        return

    if os.name == "nt":
        try:
            fd = os.open(filepath, os.O_RDWR)
            os.close(fd)
        except OSError:
            print("=" * 52)
            print("LOI: File parts.xlsx dang duoc mo boi Excel!")
            print("  -> Dong file Excel truoc, sau do chay lai script.")
            print("=" * 52)
            sys.exit(1)
    else:
        folder   = os.path.dirname(os.path.abspath(filepath))
        basename = os.path.basename(filepath)
        lockfile = os.path.join(folder, f".~lock.{basename}#")
        if os.path.exists(lockfile):
            print("=" * 52)
            print("LOI: File parts.xlsx dang duoc mo boi LibreOffice!")
            print("  -> Dong file truoc, sau do chay lai script.")
            print("=" * 52)
            sys.exit(1)

def open_excel_file(filepath):
    """Tu dong mo file Excel sau khi chay xong."""
    print(f"\n📂 Dang tu dong mo file {filepath}...")
    try:
        if platform.system() == 'Windows':
            os.startfile(filepath)
        elif platform.system() == 'Darwin':
            subprocess.call(['open', filepath])
        else:
            subprocess.call(['xdg-open', filepath])
    except Exception as e:
        print(f"⚠ Khong the tu dong mo file: {e}")

# ==============================================================================
# CAU HINH THU VIEN
# ==============================================================================
LIB_PREFIX     = "Duc"
SYMBOLS_DIR    = "symbols"
FOOTPRINTS_DIR = os.path.join("footprints", f"{LIB_PREFIX}.pretty")
MODELS_DIR     = os.path.join("3dmodels",   f"{LIB_PREFIX}.3dshapes")
TEMP_DIR       = "temp"
PARTS_FILE     = "parts.xlsx"
SHEET_LIBRARY  = "Library"
SHEET_QUEUE    = "Queue"
DEFAULT_CAT    = "Misc"

CATEGORY_MAP = {
    "resistors": "Resistors", "resistor": "Resistors", "res": "Resistors", "r": "Resistors",
    "capacitors": "Capacitors", "capacitor": "Capacitors", "cap": "Capacitors", "c": "Capacitors",
    "inductors": "Inductors", "inductor": "Inductors", "ind": "Inductors", "l": "Inductors",
    "diodes": "Diodes", "diode": "Diodes", "d": "Diodes",
    "transistors": "Transistors", "transistor": "Transistors", "mosfet": "Transistors",
    "bjt": "Transistors", "q": "Transistors",
    "mcus": "MCUs", "mcu": "MCUs", "microcontroller": "MCUs",
    "ic": "ICs", "ics": "ICs", "logic": "ICs",
    "power": "Power", "ldo": "Power", "dcdc": "Power", "regulator": "Power",
    "sensor": "Sensors", "sensors": "Sensors",
    "connectors": "Connectors", "connector": "Connectors", "conn": "Connectors",
    "crystals": "Crystals", "crystal": "Crystals", "osc": "Crystals",
    "memory": "Memory", "flash": "Memory", "eeprom": "Memory",
    "rf": "RF", "wireless": "RF", "bluetooth": "RF", "wifi": "RF",
    "analog": "Analog", "opamp": "Analog", "adc": "Analog", "dac": "Analog",
    "interface": "Interface",
    "optocouplers": "Optocouplers", "optocoupler": "Optocouplers",
    "misc": "Misc", "other": "Misc",
}

def normalize_category(raw):
    key = raw.strip().lower()
    return CATEGORY_MAP.get(key, raw.strip().title() if raw.strip() else DEFAULT_CAT)

def sym_file_for(category):
    return os.path.join(SYMBOLS_DIR, f"{LIB_PREFIX}-{category}.kicad_sym")

# ==============================================================================
# EXCEL HELPERS
# ==============================================================================
THIN = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
FONT_NORMAL = Font(name="Arial", size=10)
ALIGN_CTR   = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT  = Alignment(horizontal="left",   vertical="center")
FILL_GRAY   = PatternFill("solid", fgColor="F2F2F2")
FILL_GREEN  = PatternFill("solid", fgColor="E2EFDA")
FILL_RED    = PatternFill("solid", fgColor="FFDDE0")

def read_queue(wb):
    ws, parts = wb[SHEET_QUEUE], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        lcsc = str(row[0]).strip() if row[0] else ""
        if not lcsc or lcsc == "None":
            continue
        parts.append({
            "lcsc":     lcsc,
            "name":     str(row[1]).strip() if row[1] else "",
            "category": normalize_category(str(row[2]).strip() if row[2] else ""),
            "note":     str(row[3]).strip() if row[3] else "",
        })
    return parts

def get_library_lcsc(wb):
    ws, result = wb[SHEET_LIBRARY], set()
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0]:
            result.add(str(row[0]).strip())
    return result

def add_to_library(wb, part, sym_ok, fp_ok, model_ok):
    ws   = wb[SHEET_LIBRARY]
    nrow = ws.max_row + 1
    now  = datetime.now().strftime("%Y-%m-%d %H:%M")
    vals = [
        part["lcsc"], part["name"], part["category"], part["note"],
        "v" if sym_ok else "x",
        "v" if fp_ok  else "x",
        "v" if model_ok else "x",
        now,
    ]
    for col, val in enumerate(vals, 1):
        cell = ws.cell(row=nrow, column=col, value=val)
        cell.font   = FONT_NORMAL
        cell.border = THIN
        if col in (5, 6, 7):
            cell.alignment = ALIGN_CTR
            cell.fill = FILL_GREEN if val == "v" else FILL_RED
        elif col == 8:
            cell.alignment = ALIGN_CTR
            cell.fill = FILL_GRAY
        else:
            cell.alignment = ALIGN_LEFT
            cell.fill = FILL_GRAY

def clear_queue(wb):
    ws = wb[SHEET_QUEUE]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.value = None

def write_back_to_queue(wb, parts):
    ws = wb[SHEET_QUEUE]
    for row_idx, part in enumerate(parts, 2):
        vals = [part["lcsc"], part["name"], part["category"], part["note"]]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.font      = FONT_NORMAL
            cell.alignment = ALIGN_CTR if col == 1 else ALIGN_LEFT
            cell.border    = THIN

# ==============================================================================
# KICAD SYMBOL HELPERS
# ==============================================================================
def ensure_directories():
    for d in [SYMBOLS_DIR, FOOTPRINTS_DIR, MODELS_DIR, TEMP_DIR]:
        os.makedirs(d, exist_ok=True)

def ensure_sym_file(sym_file):
    if not os.path.exists(sym_file):
        with open(sym_file, "w", encoding="utf-8") as f:
            f.write("(kicad_symbol_lib (version 20231120) (generator easyeda2kicad)\n)\n")

def clear_temp():
    if os.path.exists(TEMP_DIR):
        shutil.rmtree(TEMP_DIR, ignore_errors=True)
    os.makedirs(TEMP_DIR, exist_ok=True)

def extract_symbol_blocks(filepath):
    """Trich xuat symbol cung cac thong tin Description de update nguoc ve Excel"""
    if not os.path.exists(filepath):
        return []
    with open(filepath, "r", encoding="utf-8") as f:
        content = f.read()
    blocks, pos = [], 0
    while True:
        start = content.find('(symbol "', pos)
        if start == -1:
            break
        stack, end = 0, -1
        for i in range(start, len(content)):
            if content[i] == "(":
                stack += 1
            elif content[i] == ")":
                stack -= 1
                if stack == 0:
                    end = i + 1
                    break
        if end == -1:
            break
        block = content[start:end]
        m = re.match(r'\(symbol\s+"([^"]+)"', block)
        if m:
            lm   = re.search(r'\(property\s+"LCSC Part"\s+"([^"]+)"', block)
            lcsc = lm.group(1) if lm else None
            
            dm   = re.search(r'\(property\s+"Description"\s+"([^"]*)"', block)
            desc = dm.group(1) if dm else ""
            
            blocks.append({"name": m.group(1), "lcsc": lcsc, "desc": desc, "content": block})
        pos = end
    return blocks

def load_all_existing():
    """Doc toan bo thu vien Kicad va lap ban do data de Auto-Fill cho Excel"""
    names = set()
    lcsc_map = {} 
    if os.path.exists(SYMBOLS_DIR):
        for fname in os.listdir(SYMBOLS_DIR):
            if fname.startswith(f"{LIB_PREFIX}-") and fname.endswith(".kicad_sym"):
                category = fname.replace(f"{LIB_PREFIX}-", "").replace(".kicad_sym", "")
                for b in extract_symbol_blocks(os.path.join(SYMBOLS_DIR, fname)):
                    names.add(b["name"])
                    if b["lcsc"]:
                        lcsc_map[b["lcsc"]] = {
                            "name": b["name"], 
                            "desc": b["desc"], 
                            "category": category
                        }
    return names, lcsc_map

def merge_symbol(block, sym_file):
    with open(sym_file, "r", encoding="utf-8") as f:
        content = f.read()
    block = re.sub(
        r'\(property\s+"Footprint"\s+"[^:]+:([^"]+)"',
        rf'(property "Footprint" "{LIB_PREFIX}:\1"',
        block
    )
    last = content.rfind(")")
    if last != -1:
        with open(sym_file, "w", encoding="utf-8") as f:
            f.write(content[:last] + "\n  " + block + "\n" + content[last:])

def process_footprints(stats):
    fp_dir  = os.path.join(TEMP_DIR, "temp.pretty")
    pattern = re.compile(r'\(model\s+"[^"]*?([^"/]+\.(?:step|wrl|STEP|WRL))"')
    added   = False
    if not os.path.exists(fp_dir):
        return False
    for fname in os.listdir(fp_dir):
        if not fname.endswith(".kicad_mod"):
            continue
        src = os.path.join(fp_dir, fname)
        dst = os.path.join(FOOTPRINTS_DIR, fname)
        if os.path.exists(dst):
            added = True
            continue
        with open(src, "r", encoding="utf-8") as f:
            data = f.read()
        data = pattern.sub(
            r'(model "${KICAD_USER_LIB}/3dmodels/' + LIB_PREFIX + r'.3dshapes/\1"',
            data
        )
        with open(dst, "w", encoding="utf-8") as f:
            f.write(data)
        print(f"  + Footprint: {fname}")
        stats["fp"] += 1
        added = True
    return added

def process_3d_models(stats):
    d3_dir = os.path.join(TEMP_DIR, "temp.3dshapes")
    added  = False
    if not os.path.exists(d3_dir):
        return False
    for fname in os.listdir(d3_dir):
        if not fname.lower().endswith((".step", ".wrl")):
            continue
        src = os.path.join(d3_dir, fname)
        dst = os.path.join(MODELS_DIR, fname)
        if os.path.exists(dst):
            added = True
            continue
        shutil.copy2(src, dst)
        print(f"  + 3D: {fname}")
        stats["3d"] += 1
        added = True
    return added

# ==============================================================================
# MAIN
# ==============================================================================
def main():
    if not os.path.exists(PARTS_FILE):
        print(f"Khong tim thay {PARTS_FILE}")
        sys.exit(1)

    check_file_not_open(PARTS_FILE)

    ensure_directories()
    wb = load_workbook(PARTS_FILE)

    if SHEET_LIBRARY not in wb.sheetnames or SHEET_QUEUE not in wb.sheetnames:
        print(f"Thieu sheet Library hoac Queue trong {PARTS_FILE}")
        sys.exit(1)

    queue = read_queue(wb)
    existing_names, existing_lcsc_map = load_all_existing()
    existing_lcsc_lib = get_library_lcsc(wb)

    already_have = []   
    to_download  = []   

    for part in queue:
        lcsc = part["lcsc"]
        lcsc_c = lcsc if lcsc.upper().startswith("C") else f"C{lcsc}"
        
        lib_info = existing_lcsc_map.get(lcsc_c) or existing_lcsc_map.get(lcsc)
        if lib_info:
            if not part["name"] or part["name"] == "None":
                part["name"] = lib_info["name"]
            if not part["note"] or part["note"] == "None":
                part["note"] = lib_info["desc"]
            if not part["category"] or part["category"] == DEFAULT_CAT:
                part["category"] = lib_info["category"]
                
            already_have.append(part)
        elif lcsc_c in existing_lcsc_lib or lcsc in existing_lcsc_lib:
            already_have.append(part)
        else:
            to_download.append(part)

    print(f"\nQueue hien tai  : {len(queue)} parts")
    print(f"Da co trong lib : {len(already_have)} - se chuyen sang Library sheet")
    print(f"Can tai moi     : {len(to_download)}\n")

    stats           = {"sym": 0, "fp": 0, "3d": 0, "success": 0, "skip": 0, "err": 0}
    processed_parts = []
    failed_parts    = []

    use_shell = (os.name == "nt")

    for i, part in enumerate(to_download, 1):
        lcsc_id  = part["lcsc"]
        note     = part["note"]
        category = part["category"]
        sym_file = sym_file_for(category)

        ensure_sym_file(sym_file)
        clear_temp()

        label = f" - {note}" if note else ""
        print(f"[{i:2}/{len(to_download)}] [{category}] {lcsc_id}{label}")

        # SỬ DỤNG TRỰC TIẾP LỆNH MẶC ĐỊNH
        result = subprocess.run(
            ["easyeda2kicad", "--full", f"--lcsc_id={lcsc_id}",
             "--output", os.path.join("temp", "temp")],
            capture_output=True, text=True, shell=use_shell
        )

        if result.returncode != 0:
            if result.stderr.strip():
                print(f"  stderr: {result.stderr.strip()[:300]}")
            print(f"  FAIL - giu lai Queue de retry")
            stats["err"] += 1
            failed_parts.append(part)
            print()
            continue

        temp_sym = os.path.join(TEMP_DIR, "temp.kicad_sym")
        blocks   = extract_symbol_blocks(temp_sym)

        if not blocks:
            print(f"  FAIL - khong co symbol hop le")
            stats["err"] += 1
            failed_parts.append(part)
            print()
            continue

        sym_ok = False
        for block in blocks:
            bname = block["name"]
            blcsc = block["lcsc"]
            bcont = block["content"]

            if bname in existing_names:
                print(f"  skip symbol '{bname}' da ton tai")
                stats["skip"] += 1
                sym_ok = True
                continue
            if blcsc and blcsc in existing_lcsc_map:
                print(f"  skip LCSC {blcsc} da ton tai")
                stats["skip"] += 1
                sym_ok = True
                continue

            if note:
                safe = note.replace('"', '\\"')
                if '(property "Description"' in bcont:
                    bcont = re.sub(
                        r'\(property\s+"Description"\s+"[^"]*"',
                        f'(property "Description" "{safe}"',
                        bcont, count=1
                    )
                else:
                    bcont = re.sub(
                        r'(\(symbol\s+"[^"]+"\n)',
                        f'\\1  (property "Description" "{safe}" (at 0 0 0)'
                        f' (effects (font (size 1.27 1.27)) (hide yes)))\n',
                        bcont, count=1
                    )

            merge_symbol(bcont, sym_file)
            existing_names.add(bname)
            if blcsc:
                existing_lcsc_map[blcsc] = {"name": bname, "desc": note, "category": category}

            print(f"  + Symbol: {bname} -> {LIB_PREFIX}-{category}")
            stats["sym"] += 1
            sym_ok = True

        fp_ok    = process_footprints(stats)
        model_ok = process_3d_models(stats)

        processed_parts.append((part, sym_ok, fp_ok, model_ok))
        stats["success"] += 1
        print()

    clear_temp()

    # ==========================================================================
    # CAP NHAT EXCEL
    # ==========================================================================
    for part, sym_ok, fp_ok, model_ok in processed_parts:
        add_to_library(wb, part, sym_ok, fp_ok, model_ok)

    for part in already_have:
        add_to_library(wb, part, True, True, True)

    # ĐỒNG BỘ CÁC Ô TRỐNG TRONG SHEET LIBRARY TỪ KICAD
    ws_lib = wb[SHEET_LIBRARY]
    for row_idx in range(2, ws_lib.max_row + 1):
        lcsc_cell = ws_lib.cell(row=row_idx, column=1)
        lcsc_val = str(lcsc_cell.value).strip() if lcsc_cell.value else ""
        lib_info = existing_lcsc_map.get(lcsc_val)
        
        if lib_info:
            name_cell = ws_lib.cell(row=row_idx, column=2)
            if not name_cell.value or str(name_cell.value).strip() == "None":
                name_cell.value = lib_info["name"]
            cat_cell = ws_lib.cell(row=row_idx, column=3)
            if not cat_cell.value or str(cat_cell.value).strip() == "None":
                cat_cell.value = lib_info["category"]
            note_cell = ws_lib.cell(row=row_idx, column=4)
            if not note_cell.value or str(note_cell.value).strip() == "None":
                note_cell.value = lib_info["desc"]

    # PHỤC HỒI LINH KIỆN MỒ CÔI TỪ KICAD VÀO SHEET LIBRARY
    current_lib_lcscs = get_library_lcsc(wb)
    recovered_count = 0
    for lcsc_kicad, info in existing_lcsc_map.items():
        if lcsc_kicad and lcsc_kicad not in current_lib_lcscs and lcsc_kicad.replace("C", "") not in current_lib_lcscs:
            missing_part = {
                "lcsc": lcsc_kicad,
                "name": info["name"],
                "category": info["category"],
                "note": info["desc"]
            }
            add_to_library(wb, missing_part, True, True, True)
            current_lib_lcscs.add(lcsc_kicad)
            recovered_count += 1
            
    if recovered_count > 0:
        print(f"  + Đã tự động phục hồi {recovered_count} linh kiện từ KiCad vào sheet Library.")

    clear_queue(wb)
    write_back_to_queue(wb, failed_parts)

    # =============== RESET GÓC NHÌN TRƯỚC KHI LƯU ===============
    for sheet in wb.worksheets:
        if sheet.sheet_view.selection:
            sheet.sheet_view.selection[0].activeCell = 'A1'
            sheet.sheet_view.selection[0].sqref = 'A1'
        sheet.sheet_view.topLeftCell = 'A1'
        
    if SHEET_QUEUE in wb.sheetnames:
        wb.active = wb.sheetnames.index(SHEET_QUEUE)
    # ===========================================================

    wb.save(PARTS_FILE)

    moved = len(processed_parts) + len(already_have) + recovered_count
    print(f"Cap nhat {PARTS_FILE}:")
    print(f"  + {moved} parts -> Library sheet")
    if failed_parts:
        print(f"  ! {len(failed_parts)} parts loi -> giu lai Queue")

    sym_files = sorted(
        f for f in os.listdir(SYMBOLS_DIR)
        if f.startswith(f"{LIB_PREFIX}-") and f.endswith(".kicad_sym")
    )
    if sym_files:
        print("\nFiles thu vien:")
        for fname in sym_files:
            blocks   = extract_symbol_blocks(os.path.join(SYMBOLS_DIR, fname))
            nickname = fname.replace(".kicad_sym", "")
            print(f"  {nickname:<28} ({len(blocks)} symbols)")

    print("\n--------------------------------")
    print(f"Symbol moi   : {stats['sym']}")
    print(f"Footprint moi: {stats['fp']}")
    print(f"3D model moi : {stats['3d']}")
    print(f"Thanh cong   : {stats['success']}")
    print(f"Bo qua (trung): {stats['skip'] + len(already_have)}")
    print(f"Loi          : {stats['err']}")

    if stats["success"] > 0:
        print("\n  git add .")
        print('  git commit -m "Update library"')
        print("  git push")

    open_excel_file(PARTS_FILE)

if __name__ == "__main__":
    main()